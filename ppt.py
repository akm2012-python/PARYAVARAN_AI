from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
import os

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

# --- Modern Color Palette ---
DARK_BG = RGBColor(15, 23, 42)      # Deep Navy / Slate
ACCENT_BLUE = RGBColor(14, 165, 233) # Cyber Blue
ACCENT_GREEN = RGBColor(34, 197, 94) # Eco Green
CARD_BG = RGBColor(255, 255, 255)    # Clean White
CARD_BORDER = RGBColor(226, 232, 240) # Slate Light Border
TEXT_DARK = RGBColor(30, 41, 59)     # Dark Slate
TEXT_LIGHT = RGBColor(100, 116, 139) # Muted Text

def apply_background(slide):
    # Top Accent Bar
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(0.15))
    bar.fill.solid()
    bar.fill.fore_color.rgb = ACCENT_GREEN
    bar.line.fill.background()

def add_header(slide, title_text, category_text="PARYAVARAN AI | SMART CITY INFRASTRUCTURE"):
    # Category Tag
    tx_cat = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.5), Inches(0.4))
    tf_cat = tx_cat.text_frame
    p_cat = tf_cat.paragraphs[0]
    p_cat.text = category_text.upper()
    p_cat.font.size = Pt(10)
    p_cat.font.bold = True
    p_cat.font.color.rgb = ACCENT_BLUE

    # Title
    tx_title = slide.shapes.add_textbox(Inches(0.8), Inches(0.7), Inches(11.5), Inches(0.8))
    tf_title = tx_title.text_frame
    p_title = tf_title.paragraphs[0]
    p_title.text = title_text
    p_title.font.size = Pt(28)
    p_title.font.bold = True
    p_title.font.color.rgb = TEXT_DARK

def create_card_grid(slide, items, columns=2):
    """Dynamically layouts items into visual card grid boxes"""
    rows = (len(items) + columns - 1) // columns
    
    start_left = Inches(0.8)
    start_top = Inches(1.6)
    total_width = Inches(11.733)
    total_height = Inches(5.2)
    
    gap = Inches(0.25)
    card_width = (total_width - (gap * (columns - 1))) / columns
    card_height = (total_height - (gap * (rows - 1))) / rows

    for i, item in enumerate(items):
        r = i // columns
        c = i % columns
        
        left = start_left + c * (card_width + gap)
        top = start_top + r * (card_height + gap)
        
        # Draw Card
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, card_width, card_height)
        card.fill.solid()
        card.fill.fore_color.rgb = CARD_BG
        card.line.color.rgb = CARD_BORDER
        card.line.width = Pt(1)
        
        # Add Card Title & Content
        tx = slide.shapes.add_textbox(left + Inches(0.15), top + Inches(0.15), card_width - Inches(0.3), card_height - Inches(0.3))
        tf = tx.text_frame
        tf.word_wrap = True
        
        if isinstance(item, tuple):
            header, desc = item
            p1 = tf.paragraphs[0]
            p1.text = header
            p1.font.bold = True
            p1.font.size = Pt(16)
            p1.font.color.rgb = ACCENT_BLUE
            p1.space_after = Pt(6)
            
            p2 = tf.add_paragraph()
            p2.text = desc
            p2.font.size = Pt(13)
            p2.font.color.rgb = TEXT_DARK
        else:
            p = tf.paragraphs[0]
            p.text = f"• {item}"
            p.font.size = Pt(14)
            p.font.color.rgb = TEXT_DARK

# --- SLIDE BUILDERS ---

# Slide 1: Modern Visual Title Slide
slide_layout = prs.slide_layouts[6]
slide1 = prs.slides.add_slide(slide_layout)

# Check if image exists
if os.path.exists("slide1_image.png"):
    slide1.shapes.add_picture("slide1_image.png", Inches(0), Inches(0), Inches(13.333), Inches(7.5))
else:
    # Banner Dark
    bg = slide1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5))
    bg.fill.solid()
    bg.fill.fore_color.rgb = DARK_BG
    
    tx = slide1.shapes.add_textbox(Inches(1.5), Inches(2.5), Inches(10.333), Inches(3))
    tf = tx.text_frame
    p1 = tf.paragraphs[0]
    p1.text = "PARYAVARAN AI"
    p1.font.bold = True
    p1.font.size = Pt(48)
    p1.font.color.rgb = ACCENT_GREEN
    
    p2 = tf.add_paragraph()
    p2.text = "AI Powered Sustainable Smart Bus Stand"
    p2.font.size = Pt(24)
    p2.font.color.rgb = RGBColor(255, 255, 255)

# Slide Data List
slides_data = [
    ("The Core Problems in Public Transit", [
        ("Wasted Electricity", "Conventional bus stands keep lights and fans running continuously even when completely empty."),
        ("Overflowing Dustbins", "Bins lack real-time sensors, leading to unhygienic conditions and delayed clearing."),
        ("Unmonitored Pollution", "Air Quality Index (AQI) and heat stress levels are ignored in public waiting spots."),
        ("Passive Infrastructure", "Existing bus stands lack intelligent decision-making, IoT connectivity, or data logging.")
    ], 2),

    ("Project Key Objectives", [
        ("Smart Energy Saving", "Reduce power consumption using passenger-detecting IR & LDR sensors."),
        ("Environmental Health", "Track Air Quality (MQ135) and temperature (DHT11) in real-time."),
        ("Automated Comfort", "Activate fans and safety lighting intelligently based on live occupancy."),
        ("Data Logging & ML", "Generate clean CSV datasets to train predictive Machine Learning models.")
    ], 2),

    ("Hardware System Architecture", [
        ("Microcontroller", "Arduino UNO processes sensor inputs and drives relays."),
        ("Environment Sensors", "DHT11 (Temp/Humidity) & MQ135 (Air Quality)."),
        ("Occupancy & Range", "IR Sensor (Passenger presence) & Ultrasonic HC-SR04 (Dustbin level)."),
        ("Actuators & Controls", "5V Relays, Cooling Fan, Safety Lights, Buzzer & Traffic Status Module.")
    ], 2),

    ("Explainable AI (XAI) Dashboard Logic", [
        ("Rule-Based Intelligence", "System evaluates multi-sensor thresholds before taking action."),
        ("XAI Transparency", "Dashboard explains WHY decisions were made (e.g., Fan ON because Temp > 30°C AND Passenger Present)."),
        ("Confidence Score", "Outputs confidence percentages for system recommendations (e.g., 92% confidence)."),
        ("Operator Trust", "Eliminates black-box automation by giving clear reasoning to city operators.")
    ], 2),

    ("Ternary Environmental Logic", [
        ("State +1 (Green)", "Optimal Environment: Clean air, normal temp, low energy state."),
        ("State 0 (Yellow)", "Moderate Caution: Raised temperature or moderate AQI warning."),
        ("State -1 (Red)", "Critical Alert: High pollution detected or overflowing dustbin level."),
        ("Better Than Binary", "Provides 3 intuitive status conditions instead of simple ON/OFF binary states.")
    ], 3),

    ("Dataset Pipeline & Research Roadmap", [
        ("Automated Logging", "Logs sensor values, occupancy, and AI scores every 2 seconds into CSV."),
        ("Model Evaluation", "Compares Decision Trees, Random Forest, XGBoost, SVM, and KNN."),
        ("Accuracy Optimization", "Determines which algorithm predicts energy saving potential most accurately."),
        ("Academic Goal", "Paves the way for a published research paper on AI smart infrastructure.")
    ], 2),

    ("Expected Impact & Results", [
        ("Energy Saved", "~42% reduction in electricity waste recorded during testing."),
        ("Instant Response", "Automated lighting and fan response within 500ms of passenger detection."),
        ("Predictive Waste Mgmt", "Alerts authorities before dustbins reach 90% capacity."),
        ("Smart City Ready", "Low-cost model easily scalable using ESP32 and TinyML.")
    ], 2)
]

# Generate Grid Slides
for title, items, cols in slides_data:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    apply_background(slide)
    add_header(slide, title)
    create_card_grid(slide, items, columns=cols)

prs.save("PARYAVARAN_AI_Modern_Presentation.pptx")
print("Successfully generated modern deck: PARYAVARAN_AI_Modern_Presentation.pptx")